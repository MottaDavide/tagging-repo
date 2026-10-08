"""Costruzione dell'Excel di sintesi delle metriche (per documento e per Risk Model).

Unica rappresentazione visiva dei risultati calcolati da `evaluation.services.matcher`:
niente logica di matching qui, solo rendering di `DocResult`/`Counts` già calcolati.
"""

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from evaluation.models.results import Counts, DocResult
from evaluation.services.matcher import doc_rows, rm_rows

_FONT = "Calibri"
_HAIRLINE = "ECEFF4"
_BODY_TEXT = "3A4A63"

# Tema pastello per foglio: fill header, testo header, riga sotto header,
# tinta bande alternate, fill riga totali, riga sopra i totali.
_THEME_DOC = {
    "fill": "DCE3F2",
    "text": "33415C",
    "rule": "B7C5E0",
    "band": "F7F9FC",
    "total": "EEF2FA",
    "accent": "9DB2D9",
}
_THEME_RM = {
    "fill": "DCEAD9",
    "text": "41513C",
    "rule": "BAD3B4",
    "band": "F6FAF4",
    "total": "EAF2E7",
    "accent": "A9C99E",
}

_DOC_HEADERS = ["Documento", "GT", "PRED", "TP", "FP", "FN", "Recall", "Precisione"]
_RM_HEADERS = ["Risk Model", "TP", "FP", "FN", "Precisione", "Recall", "F1"]

Row = list[str | int | float]


def _subtitle(rm_filter: frozenset[str], tipo_filter: frozenset[str], ignore_ente: bool, threshold: float) -> str:
    tuple_label = "(RM, MO, OdC)" if ignore_ente else "(RM, Ente, MO, OdC)"
    tipo_label = ", ".join(sorted(tipo_filter)) if tipo_filter else "tutti"
    return (
        f"Filtro RM: {', '.join(sorted(rm_filter))}   ·   Filtro Tipo Citazione: {tipo_label}   ·   "
        f"Matching: n-upla {tuple_label} esatta, pagine sovrapposte, overlap citazione ≥ {threshold:.0%}"
    )


def _write_title(ws: Worksheet, title: str, subtitle: str, n_cols: int, theme: dict[str, str]) -> int:
    """Scrive titolo + sottotitolo nelle prime righe e ritorna la riga dell'header tabella."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    t = ws.cell(row=1, column=1, value=title)
    t.font = Font(name=_FONT, size=16, bold=True, color=theme["text"])
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    s = ws.cell(row=2, column=1, value=subtitle)
    s.font = Font(name=_FONT, size=10, italic=True, color="8A93A6")
    ws.row_dimensions[2].height = 18

    accent = Border(bottom=Side(style="medium", color=theme["accent"]))
    for col in range(1, n_cols + 1):
        ws.cell(row=2, column=col).border = accent
    return 4


def _style_table(
    ws: Worksheet,
    header_row: int,
    n_data_rows: int,
    n_cols: int,
    theme: dict[str, str],
    totals_row: bool = False,
) -> None:
    """Header pastello, bande alternate, hairline orizzontali.

    Se `totals_row`, l'ultima riga è trattata come riga di totale (grassetto + fill dedicato) —
    da passare solo quando quella riga è davvero un aggregato, non l'ultimo elemento della lista.
    """
    header_fill = PatternFill("solid", fgColor=theme["fill"])
    band_fill = PatternFill("solid", fgColor=theme["band"])
    total_fill = PatternFill("solid", fgColor=theme["total"])
    header_border = Border(bottom=Side(style="thin", color=theme["rule"]))
    total_border = Border(top=Side(style="medium", color=theme["accent"]))
    hairline = Border(bottom=Side(style="thin", color=_HAIRLINE))
    end_row = header_row + n_data_rows

    for row in range(header_row, end_row + 1):
        is_header = row == header_row
        is_total = totals_row and row == end_row and not is_header
        is_band = not is_header and not is_total and (row - header_row) % 2 == 0
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=col)
            if is_header:
                cell.font = Font(name=_FONT, size=11, bold=True, color=theme["text"])
                cell.fill = header_fill
                cell.border = header_border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                continue
            cell.font = Font(name=_FONT, size=11, bold=is_total, color=_BODY_TEXT)
            cell.border = total_border if is_total else hairline
            if is_total:
                cell.fill = total_fill
            elif is_band:
                cell.fill = band_fill
            cell.alignment = Alignment(horizontal="left" if col == 1 else "right", vertical="center")

    ws.row_dimensions[header_row].height = 26
    for row in range(header_row + 1, end_row + 1):
        ws.row_dimensions[row].height = 20
    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False


def _autofit(ws: Worksheet, min_width: int = 10, max_width: int = 50) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = max(len(line) for line in str(cell.value).split("\n"))
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, length in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, length + 2))


def _write_metrics_sheet(
    ws: Worksheet,
    title: str,
    subtitle: str,
    headers: list[str],
    rows: list[Row],
    n_pct_cols: int,
    theme: dict[str, str],
    totals_row: bool = False,
) -> None:
    """Scrive titolo + tabella su `ws`. Le ultime `n_pct_cols` colonne sono percentuali,
    le colonne numeriche restanti (dalla 2ª in poi) sono formattate come interi."""
    header_row = _write_title(ws, title, subtitle, len(headers), theme)
    for col, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=h)

    n_cols = len(headers)
    pct_start = n_cols - n_pct_cols + 1
    for i, values in enumerate(rows):
        row = header_row + 1 + i
        for col, v in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=v)
        for col in range(2, pct_start):
            ws.cell(row=row, column=col).number_format = "#,##0"
        for col in range(pct_start, n_cols + 1):
            ws.cell(row=row, column=col).number_format = "0.0%"

    _style_table(ws, header_row, len(rows), n_cols, theme, totals_row=totals_row)
    _autofit(ws)


def _write_per_doc_sheet(wb: openpyxl.Workbook, docs: list[DocResult], subtitle: str) -> None:
    ws = wb.active
    assert ws is not None
    ws.title = "Per documento"

    rows: list[Row] = [
        [d.doc, len(d.gt), len(d.pred), m.counts.tp, m.counts.fp, m.counts.fn, m.recall, m.precision]
        for d, m in doc_rows(docs)
    ]
    _write_metrics_sheet(ws, "METRICHE DI TAGGING — per documento", subtitle, _DOC_HEADERS, rows, 2, _THEME_DOC)


def _write_per_rm_sheet(wb: openpyxl.Workbook, per_rm: dict[str, Counts], total: Counts, subtitle: str) -> None:
    ws = wb.create_sheet("Per Risk Model")

    rows: list[Row] = [
        [label, m.counts.tp, m.counts.fp, m.counts.fn, m.precision, m.recall, m.f1]
        for label, m in rm_rows(per_rm, total)
    ]
    _write_metrics_sheet(
        ws, "METRICHE DI TAGGING — per Risk Model", subtitle, _RM_HEADERS, rows, 3, _THEME_RM, totals_row=True
    )


def write_summary(
    docs: list[DocResult],
    per_rm: dict[str, Counts],
    total: Counts,
    *,
    rm_filter: frozenset[str],
    tipo_filter: frozenset[str],
    ignore_ente: bool,
    threshold: float,
    output: Path,
) -> None:
    """Salva in `output` l'Excel di sintesi: fogli 'Per documento' e 'Per Risk Model'."""
    subtitle = _subtitle(rm_filter, tipo_filter, ignore_ente, threshold)
    wb = openpyxl.Workbook()
    _write_per_doc_sheet(wb, docs, subtitle)
    _write_per_rm_sheet(wb, per_rm, total, subtitle)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
